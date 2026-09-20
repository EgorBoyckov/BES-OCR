import os

from docx import Document as DocxDocument

from bes_ocr.config.settings import Settings
from bes_ocr.core.docx_writer import build_docx
from bes_ocr.core.errors import PdfOpenError
from bes_ocr.core.pipeline import process_document
import pytest


def _convert(pdf_path, tmp_path, settings=None):
    settings = settings or Settings()
    document = process_document(pdf_path, settings)
    out_path = os.path.join(tmp_path, "out.docx")
    build_docx(document, out_path)
    return document, out_path


def test_text_pdf_end_to_end(text_pdf, tmp_path):
    document, out_path = _convert(text_pdf, tmp_path)
    assert os.path.isfile(out_path)
    reopened = DocxDocument(out_path)
    assert len(reopened.tables) == 1
    assert not document.warnings


def test_scanned_russian_pdf_end_to_end(russian_scanned_pdf, tmp_path):
    document, out_path = _convert(russian_scanned_pdf, tmp_path)
    assert document.pages[0].used_ocr
    reopened = DocxDocument(out_path)
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "русском" in full_text or "документ" in full_text
    assert len(reopened.tables) == 1


def test_english_pdf_end_to_end(english_pdf, tmp_path):
    document, out_path = _convert(english_pdf, tmp_path)
    reopened = DocxDocument(out_path)
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "English" in full_text


def test_mixed_language_pdf(mixed_language_pdf, tmp_path):
    document, out_path = _convert(mixed_language_pdf, tmp_path)
    assert os.path.isfile(out_path)


def test_table_with_merged_cells(table_merged_cells_pdf, tmp_path):
    document, out_path = _convert(table_merged_cells_pdf, tmp_path)
    reopened = DocxDocument(out_path)
    assert len(reopened.tables) == 1
    t = reopened.tables[0]
    assert t.cell(0, 0).text == t.cell(0, 1).text


def test_multipage_table_pdf(multipage_table_pdf, tmp_path):
    document, out_path = _convert(multipage_table_pdf, tmp_path)
    reopened = DocxDocument(out_path)
    assert len(reopened.tables) == 2


def test_images_pdf(images_pdf, tmp_path):
    document, out_path = _convert(images_pdf, tmp_path)
    reopened = DocxDocument(out_path)
    assert len(reopened.inline_shapes) >= 1


def test_complex_mixed_page(complex_mixed_pdf, tmp_path):
    document, out_path = _convert(complex_mixed_pdf, tmp_path)
    reopened = DocxDocument(out_path)
    assert len(reopened.tables) == 1
    styles = {p.style.name for p in reopened.paragraphs}
    assert any(s.startswith("List") for s in styles)


def test_multipage_document(multipage_pdf, tmp_path):
    document, out_path = _convert(multipage_pdf, tmp_path)
    assert len(document.pages) == 5


def test_large_document_100_plus_pages(large_pdf, tmp_path):
    settings = Settings()
    document, out_path = _convert(large_pdf, tmp_path, settings)
    assert len(document.pages) == 105
    assert os.path.isfile(out_path)


def test_corrupted_pdf_raises(corrupted_pdf):
    with pytest.raises(PdfOpenError):
        process_document(corrupted_pdf, Settings())


def test_no_text_layer_pdf_uses_ocr(no_text_layer_pdf, tmp_path):
    document, _ = _convert(no_text_layer_pdf, tmp_path)
    assert document.pages[0].used_ocr


def test_missing_file_raises(tmp_path):
    with pytest.raises(PdfOpenError):
        process_document(str(tmp_path / "does_not_exist.pdf"), Settings())
