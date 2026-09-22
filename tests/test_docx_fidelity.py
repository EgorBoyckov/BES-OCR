"""Генерация DOCX, воспроизводящая геометрию оригинала (core/docx_writer.py)."""
import io
import os

from docx import Document as DocxDocument
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from PIL import Image

from bes_ocr.core.docx_writer import build_docx, clean_font_name
from bes_ocr.core.models import (
    Alignment,
    BBox,
    Document,
    Heading,
    ImageBlock,
    Page,
    Paragraph,
    Run,
    Table,
    TableCell,
    TableRow,
)


def _para(text, baseline, x0=60.0, x1=400.0, size=12.0, pitch=None, n_lines=1, **kw):
    return Paragraph(
        runs=[Run(text=text, size_pt=size)],
        bbox=BBox(x0, baseline - 0.7 * size, x1, baseline + 0.2 * size + (n_lines - 1) * (pitch or 0)),
        indent_pt=x0,
        baseline_pt=baseline,
        line_pitch_pt=pitch,
        n_lines=n_lines,
        **kw,
    )


def _save(document, tmp_path):
    out = os.path.join(tmp_path, "out.docx")
    build_docx(document, out)
    return DocxDocument(out)


def test_exact_line_spacing_and_space_before_follow_source_baselines(tmp_path):
    page = Page(number=1)
    page.blocks.append(_para("Первый абзац", baseline=100.0, pitch=14.0, n_lines=2))
    page.blocks.append(_para("Второй абзац", baseline=160.0, pitch=14.0, n_lines=2))
    reopened = _save(Document(pages=[page]), tmp_path)
    p1, p2 = [p for p in reopened.paragraphs if p.text]
    assert p1.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert p1.paragraph_format.line_spacing.pt == 14.0
    # Второй абзац: базовая линия на 60pt ниже первой, первый занимает две
    # строки по 14pt → между абзацами 60 - 28 = 32pt.
    assert abs(p2.paragraph_format.space_before.pt - 32.0) < 0.1


def test_paragraph_after_table_is_positioned_from_table_bottom(tmp_path):
    """Регрессия: зазор перед абзацем после таблицы считался от предыдущего
    ТЕКСТОВОГО блока (над таблицей) — в DOCX после таблицы появлялась
    пустота высотой в саму таблицу."""
    page = Page(number=1)
    page.blocks.append(_para("Над таблицей", baseline=100.0))
    table = Table(
        rows=[TableRow(cells=[TableCell(blocks=[Paragraph(runs=[Run(text="A", size_pt=9)])])], height_pt=100.0)],
        col_widths_pt=[200.0],
        bbox=BBox(60.0, 120.0, 260.0, 220.0),
    )
    page.blocks.append(table)
    page.blocks.append(_para("Под таблицей", baseline=240.0))
    reopened = _save(Document(pages=[page]), tmp_path)
    after = next(p for p in reopened.paragraphs if p.text == "Под таблицей")
    assert after.paragraph_format.space_before.pt < 20.0
    before = next(p for p in reopened.paragraphs if p.text == "Над таблицей")
    assert before.paragraph_format.space_after.pt > 5.0  # зазор перед таблицей
    assert reopened.tables[0].rows[0].height.pt == 100.0


def test_forced_line_break_written_as_break_and_alignment_kept(tmp_path):
    page = Page(number=1)
    page.blocks.append(
        _para("Приложение 1\nк договору № 1", baseline=40.0, pitch=14.0, n_lines=2, alignment=Alignment.RIGHT)
    )
    reopened = _save(Document(pages=[page]), tmp_path)
    p = next(p for p in reopened.paragraphs if p.text)
    assert p.alignment == 2  # RIGHT
    assert p._p.xpath(".//w:br")


def test_fonts_are_times_by_default_and_headings_not_restyled(tmp_path):
    page = Page(number=1)
    page.blocks.append(Heading(runs=[Run(text="Заголовок", size_pt=14, bold=True)], level=1, baseline_pt=50.0, bbox=BBox(60, 40, 200, 55), indent_pt=60))
    page.blocks.append(_para("Текст", baseline=100.0))
    reopened = _save(Document(pages=[page]), tmp_path)
    heading = next(p for p in reopened.paragraphs if p.text == "Заголовок")
    run = heading.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size.pt == 14 and run.bold
    # Стиль заголовка шаблона больше не красит текст в синий и не меняет шрифт.
    style_color = reopened.styles["Heading 1"].font.color.rgb
    assert style_color is None or str(style_color) == "000000"


def test_text_layer_font_names_are_cleaned():
    assert clean_font_name("ABCDEF+TimesNewRomanPS-BoldMT") == "Times New Roman"
    assert clean_font_name("Arial,Bold") == "Arial"
    assert clean_font_name("PTAstraSerif-Regular") == "PT Astra Serif"
    assert clean_font_name(None) is None


def test_floating_image_is_anchored_at_source_position(tmp_path):
    img = Image.new("RGBA", (40, 40), (0, 0, 255, 128))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    page = Page(number=1)
    page.blocks.append(_para("Текст", baseline=100.0))
    page.blocks.append(ImageBlock(data=buf.getvalue(), width_px=40, height_px=40, bbox=BBox(300, 500, 400, 600), floating=True))
    reopened = _save(Document(pages=[page]), tmp_path)
    anchors = reopened.element.body.xpath(".//wp:anchor")
    assert len(anchors) == 1
    anchor = anchors[0]
    assert anchor.get("behindDoc") == "1"
    pos_h = anchor.find(qn("wp:positionH")).find(qn("wp:posOffset")).text
    pos_v = anchor.find(qn("wp:positionV")).find(qn("wp:posOffset")).text
    assert int(pos_h) == 300 * 12700 and int(pos_v) == 500 * 12700
    assert len(reopened.inline_shapes) == 0  # не в потоке текста


def test_borderless_layout_table_has_no_grid_style(tmp_path):
    cells = [
        TableCell(blocks=[_para("Левая колонка", baseline=100.0, x0=60, x1=300)], bbox=BBox(60, 90, 310, 150)),
        TableCell(blocks=[_para("Правая колонка", baseline=100.0, x0=320, x1=560)], bbox=BBox(310, 90, 560, 150)),
    ]
    table = Table(rows=[TableRow(cells=cells)], col_widths_pt=[250, 250], bbox=BBox(60, 90, 560, 150), borderless=True, source="layout")
    page = Page(number=1, blocks=[table])
    reopened = _save(Document(pages=[page]), tmp_path)
    t = reopened.tables[0]
    assert t.style is None or t.style.name != "Table Grid"
    assert t.cell(0, 0).text == "Левая колонка"
    assert t.cell(0, 1).text == "Правая колонка"


def test_pages_are_separated_by_page_break_property_not_extra_paragraph(tmp_path):
    pages = [Page(number=1, blocks=[_para("Страница один", baseline=100.0)]), Page(number=2, blocks=[_para("Страница два", baseline=100.0)])]
    reopened = _save(Document(pages=pages), tmp_path)
    texts = [p for p in reopened.paragraphs]
    assert [p.text for p in texts] == ["Страница один", "Страница два"]
    assert texts[1].paragraph_format.page_break_before
