"""Регрессии, найденные на реальном скане пользователя (2 страницы, 150 dpi,
синие печати, 5 таблиц): запасной детектор таблиц, нумерация списков,
накопление вертикального сдвига, полужирный, мелкие ошибки OCR."""
import os

import numpy as np
from docx import Document as DocxDocument
from PIL import Image, ImageDraw

from bes_ocr.core.docx_writer import build_docx
from bes_ocr.core.layout_analysis import _bold_threshold, build_blocks
from bes_ocr.core.models import BBox, Document, ListItem, Page, Run
from bes_ocr.core.ocr_engine import is_plausible_word
from bes_ocr.core.page_processor import _fix_vertical_bars
from bes_ocr.core.table_detection import add_missed_tables, detect_tables_ruling_lines

from .layout_helpers import justified_line, line_at, word_at


def _grid_image(wavy: bool = False) -> np.ndarray:
    """Таблица 3×3 с объединённой шапкой, бледными и слегка волнистыми
    линиями, как на реальном скане."""
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    gray = (120, 120, 120)
    xs = [100, 500, 900, 1400]
    ys = [100, 250, 400, 550]
    for i, y in enumerate(ys):
        dy = 6 if (wavy and i == 2) else 0
        d.line((xs[0], y, 700, y + dy), fill=gray, width=3)
        d.line((700, y + dy, xs[-1], y), fill=gray, width=3)
    for x in xs:
        d.line((x, ys[0], x, ys[-1]), fill=gray, width=3)
    # Шапка объединена на всю ширину: внутренние вертикали начинаются ниже.
    for x in xs[1:-1]:
        d.rectangle((x - 2, ys[0] + 3, x + 2, ys[1] - 3), fill="white")
    # Буквы с длинными вертикальными штрихами внутри ячеек не должны
    # становиться границами.
    for x in (200, 600, 1000):
        d.line((x, 300, x, 360), fill="black", width=4)
    return np.asarray(img)[:, :, ::-1].copy()


def test_ruling_line_detector_finds_grid_and_merged_header():
    tables = detect_tables_ruling_lines(_grid_image(wavy=True), px_to_pt=72 / 300)
    assert len(tables) == 1
    t = tables[0]
    assert (t.n_rows, t.n_cols) == (3, 3)
    assert t.rows[0].cells[0].col_span == 3
    assert all(c.is_merge_continuation for c in t.rows[0].cells[1:])
    assert all(not c.is_merge_continuation and c.col_span == 1 for r in t.rows[1:] for c in r.cells)


def test_missed_table_is_added_but_duplicates_are_not():
    found = detect_tables_ruling_lines(_grid_image(), px_to_pt=72 / 300)
    assert len(add_missed_tables([], found)) == 1
    assert len(add_missed_tables(found, found)) == 1


def test_list_markers_are_literal_and_numbering_restarts(tmp_path):
    """Word-автонумерация продолжала нумерацию через весь документ ("5.",
    "6." вместо "1.", "2." во втором списке) и превращала "-" в "•"."""
    page = Page(number=1)
    for i, (marker, text) in enumerate([("1.", "Первый"), ("2.", "Второй"), ("1.", "Снова первый"), ("-", "Пункт")]):
        page.blocks.append(
            ListItem(
                runs=[Run(text=text, size_pt=12)], marker=marker, ordered=marker != "-", baseline_pt=100 + i * 14,
                bbox=BBox(60, 90 + i * 14, 300, 103 + i * 14), indent_pt=75, first_line_indent_pt=-15, text_x_pt=75,
            )
        )
    out = os.path.join(tmp_path, "l.docx")
    build_docx(Document(pages=[page]), out)
    paragraphs = [p for p in DocxDocument(out).paragraphs if p.text.strip()]
    assert [p.text.split()[0] for p in paragraphs] == ["1.", "2.", "1.", "-"]
    for p in paragraphs:
        assert p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val") == ["0"]


def test_numbered_paragraph_with_continuation_at_left_margin():
    """"3. Проведена оценка ..." — номер в красной строке, продолжение от
    левого поля (не под текстом пункта)."""
    words = justified_line("3. Проведена оценка условий труда на рабочих местах используемых для", 90, 560, 100)
    words += line_at("практики:", 60, 114)
    words += justified_line("Обычный абзац текста во всю ширину страницы для определения полей", 60, 560, 140)
    item = build_blocks(words, 595, 842)[0]
    assert isinstance(item, ListItem) and item.marker == "3."
    assert item.text.endswith("практики:")
    assert item.indent_pt == 60 and item.first_line_indent_pt == 30


def test_adjacent_paragraphs_get_pitch_leading_to_next_baseline():
    """Шаг внутри абзаца (14.2) и расстояние до следующего абзаца (13.9) в
    оригинале немного различаются; отрицательного отступа в Word нет, и
    без поправки сдвиг вниз копился от абзаца к абзацу."""
    words = []
    base = 100.0
    for n_lines, pitch in ((2, 14.2), (3, 14.2), (2, 13.9)):
        for k in range(n_lines):
            text = "начало абзаца" if k == 0 else "продолжение текста абзаца"
            words += line_at(("- " if k == 0 else "") + text + " " + "слово " * 12, 60 if k else 80, base + k * pitch)
        base += n_lines * pitch - 0.3
    blocks = build_blocks(words, 595, 842, detect_columns=False)
    for a, b in zip(blocks, blocks[1:]):
        assert abs(a.baseline_pt + a.line_pitch_pt * a.n_lines - b.baseline_pt) < 0.01


def test_bold_threshold_adapts_to_scan_quality():
    # Реальный скан 150 dpi: обычный текст 0.072–0.081, полужирный ~0.10.
    real = [(0.075, 200), (0.078, 300), (0.081, 100), (0.100, 150), (0.106, 60)]
    t = _bold_threshold(real)
    assert 0.081 < t < 0.100
    # Размытый скан: классы ближе (0.088–0.092 и 0.112–0.120).
    blurry = [(0.088, 300), (0.090, 200), (0.092, 150), (0.112, 40), (0.117, 200)]
    t = _bold_threshold(blurry)
    assert 0.092 < t < 0.112
    # Только обычный текст с шумом — полужирных нет.
    plain = [(0.074, 200), (0.077, 300), (0.080, 200), (0.083, 50)]
    assert _bold_threshold(plain) > 0.083


def test_vertical_bar_of_digit_height_is_one():
    words = [word_at("Приложение", 400, 40, 12), word_at("|", 480, 40, 12)]
    words[1].y0, words[1].y1 = 40 - 8.2, 40  # высота цифры
    fixed = _fix_vertical_bars(words)
    assert [w.text for w in fixed] == ["Приложение", "1"]
    tall = word_at("|", 480, 40, 12)
    tall.y0, tall.y1 = 20, 45  # выше строки — обрывок линии
    assert [w.text for w in _fix_vertical_bars([tall])] == []


def test_plausible_words_include_emails_and_numero_sign():
    for text in ("ekb@1cbit.ru", "usue@usue.ru", "договору№", "www.usue.ru"):
        assert is_plausible_word(text), text
    assert not is_plausible_word("MG")
