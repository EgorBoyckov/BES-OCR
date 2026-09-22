"""Генерация .docx из внутренней модели Document.

Не знает ничего о PDF или OCR — только о модели `core.models`. Ключевое
требование проекта: таблицы должны быть настоящими редактируемыми таблицами
Word (`table.cell(...).merge(...)`), а не изображениями.

Для визуального сходства с оригиналом (п.20 ТЗ) страница DOCX должна иметь
тот же размер и поля, что и исходный PDF — иначе весь перенос строк,
пропорции таблиц и отступы систематически расходятся с оригиналом, даже
если каждый блок распознан правильно по отдельности. См. `_compute_page_geometry`.

Вертикальная вёрстка воспроизводится "потоком с курсором" (`_Flow`): для
каждого блока известна абсолютная позиция его первой базовой линии на
исходной странице, и отступ перед блоком вычисляется так, чтобы базовая
линия в Word оказалась на той же высоте — с учётом того, где закончился
предыдущий блок (в том числе таблица), а не как "зазор между рамками
соседних текстовых блоков" (такой зазор не учитывал таблицу между ними и
давал огромные пустоты после таблиц).
"""
from __future__ import annotations

import copy
import io
import logging
import re
import statistics
from typing import Optional

from docx import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from PIL import Image as PILImage

from .errors import DocxWriteError
from .models import (
    Alignment,
    Document,
    Heading,
    ImageBlock,
    ListItem,
    Paragraph,
    Table,
)

logger = logging.getLogger("bes_ocr")

_ALIGN_MAP = {
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_VALIGN_MAP = {
    "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
    "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
}

DEFAULT_FONT_NAME = "Times New Roman"

# Запасные поля (ГОСТ 7.32 — стандартные для русскоязычных официальных
# документов), используются только когда их не удалось оценить по
# содержимому конкретного документа.
_FALLBACK_MARGIN_LEFT_PT = 85.0
_FALLBACK_MARGIN_RIGHT_PT = 28.0
_FALLBACK_MARGIN_TOP_PT = 57.0
_FALLBACK_MARGIN_BOTTOM_PT = 57.0
_DEFAULT_PAGE_WIDTH_PT = 595.0
_DEFAULT_PAGE_HEIGHT_PT = 842.0

# Положение базовой линии внутри строки с точным межстрочным интервалом L:
# текстовый процессор распределяет высоту строки между надстрочной и
# подстрочной частями пропорционально метрикам шрифта (Times New Roman:
# ascent 0.891, descent 0.216 em), т.е. базовая линия находится на
# 0.891 / 1.107 ≈ 0.805·L от верха строки. Измерено на LibreOffice
# (базовые линии в PDF-рендере при разных L и кеглях совпадают с моделью
# до 0.1pt).
_BASELINE_RATIO = 0.891 / (0.891 + 0.216)
# Межстрочный шаг одинарного интервала для Times New Roman (доля кегля) —
# для однострочных абзацев, у которых шаг не измерить.
_DEFAULT_PITCH_RATIO = 1.15
# Внутренние поля ячеек таблиц по горизонтали, если их не удалось измерить
# по оригиналу (значение Word по умолчанию, 0.19 см).
_CELL_MARGIN_PT = 5.4
# Небольшой запас справа: если шрифт в Word окажется чуть шире исходного,
# строка не перенесётся раньше, чем в оригинале.
_WRAP_SLACK_PT = 1.5

_FONT_SUFFIX_RE = re.compile(
    r"([,-]?(Bold|Italic|Oblique|Regular|Roman|Book|Medium|Light|Semibold|Demi|BoldItalic|BoldOblique)"
    r"|PSMT|PS|MT|-Identity-H)+$",
    re.IGNORECASE,
)
_KNOWN_FONTS = {
    "timesnewroman": "Times New Roman",
    "times": "Times New Roman",
    "arial": "Arial",
    "helvetica": "Arial",
    "couriernew": "Courier New",
    "courier": "Courier New",
    "calibri": "Calibri",
    "cambria": "Cambria",
    "verdana": "Verdana",
    "tahoma": "Tahoma",
    "georgia": "Georgia",
    "liberationserif": "Liberation Serif",
    "liberationsans": "Liberation Sans",
    "dejavuserif": "DejaVu Serif",
    "dejavusans": "DejaVu Sans",
    "ptastraserif": "PT Astra Serif",
    "ptastrasans": "PT Astra Sans",
    "ptserif": "PT Serif",
    "ptsans": "PT Sans",
}


def clean_font_name(raw: Optional[str]) -> Optional[str]:
    """Имя шрифта из PDF ("ABCDEF+TimesNewRomanPS-BoldMT") → имя
    семейства для Word ("Times New Roman")."""
    if not raw:
        return None
    name = raw.split("+", 1)[-1]
    full_key = re.sub(r"[\s_,-]", "", name).lower()
    # Сначала — известные семейства по префиксу (у "TimesNewRoman" слово
    # "Roman" — часть имени, а не начертание).
    for key in sorted(_KNOWN_FONTS, key=len, reverse=True):
        if full_key.startswith(key):
            return _KNOWN_FONTS[key]
    name = _FONT_SUFFIX_RE.sub("", name).strip(" ,-")
    if not name:
        return None
    # "TimesNewRoman" → "Times New Roman"
    spaced = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name)
    return spaced


def _add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = run._r.makeelement(qn("w:instrText"), {})
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_end = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


class PageGeometry:
    def __init__(self, width_pt: float, height_pt: float, margin_left_pt: float, margin_right_pt: float,
                 margin_top_pt: float, margin_bottom_pt: float):
        self.width_pt = width_pt
        self.height_pt = height_pt
        self.margin_left_pt = margin_left_pt
        self.margin_right_pt = margin_right_pt
        self.margin_top_pt = margin_top_pt
        self.margin_bottom_pt = margin_bottom_pt

    @property
    def usable_width_pt(self) -> float:
        return max(72.0, self.width_pt - self.margin_left_pt - self.margin_right_pt)


def _percentile(values: list[float], p: float) -> float:
    if not values:
        return 0.0
    s = sorted(values)
    idx = min(len(s) - 1, max(0, int(round(p * (len(s) - 1)))))
    return s[idx]


def _dominant_size(block: Paragraph) -> float:
    pairs = [(r.size_pt, len(r.text)) for r in block.runs if r.text.strip()]
    if not pairs:
        return 11.0
    pairs.sort()
    total = sum(w for _, w in pairs)
    acc = 0
    for v, w in pairs:
        acc += w
        if acc >= total / 2:
            return v
    return pairs[-1][0]


def _block_pitch(block: Paragraph, pitch_ratio: float) -> float:
    size = _dominant_size(block)
    pitch = block.line_pitch_pt or size * pitch_ratio
    return max(pitch, size * 1.0)


def _block_box_top(block: Paragraph, pitch_ratio: float) -> Optional[float]:
    """Верх строки (line box) первой строки абзаца в координатах страницы."""
    if block.baseline_pt is None:
        return None
    return block.baseline_pt - _block_pitch(block, pitch_ratio) * _BASELINE_RATIO


def _compute_page_geometry(document: Document, pitch_ratio: float = _DEFAULT_PITCH_RATIO) -> PageGeometry:
    """Определяет размер страницы и поля DOCX-документа по исходному PDF.

    Без этого DOCX получает поля/размер страницы по умолчанию (обычно
    Letter, 1 дюйм со всех сторон) независимо от реального документа — из-за
    этого систематически расходятся перенос строк, пропорции таблиц и
    отступы, даже если каждый блок распознан правильно по отдельности.
    Поля оцениваются по фактическому содержимому (10-й/90-й процентиль
    левых/правых краёв текстовых блоков), а не берутся "как есть" —
    единичный выброс (например, широкая таблица у самого края) не должен
    задавать поля для всей страницы.
    """
    if not document.pages:
        return PageGeometry(
            _DEFAULT_PAGE_WIDTH_PT, _DEFAULT_PAGE_HEIGHT_PT,
            _FALLBACK_MARGIN_LEFT_PT, _FALLBACK_MARGIN_RIGHT_PT,
            _FALLBACK_MARGIN_TOP_PT, _FALLBACK_MARGIN_BOTTOM_PT,
        )

    width_pt = statistics.median(p.width_pt for p in document.pages) or _DEFAULT_PAGE_WIDTH_PT
    height_pt = statistics.median(p.height_pt for p in document.pages) or _DEFAULT_PAGE_HEIGHT_PT

    lefts: list[float] = []
    rights: list[float] = []
    tops: list[float] = []
    bottoms: list[float] = []
    for page in document.pages:
        text_blocks = [b for b in page.blocks if isinstance(b, (Paragraph, Heading, ListItem)) and b.bbox]
        flow_blocks = [
            b for b in page.blocks
            if getattr(b, "bbox", None) is not None and not (isinstance(b, ImageBlock) and b.floating)
        ]
        if flow_blocks:
            page_tops = []
            for b in flow_blocks:
                box_top = _block_box_top(b, pitch_ratio) if isinstance(b, Paragraph) else None
                page_tops.append(box_top if box_top is not None else b.bbox.y0)
            tops.append(min(page_tops))
            bottoms.append(max(b.bbox.y1 for b in flow_blocks))
        for b in text_blocks:
            lefts.append(b.bbox.x0)
            rights.append(b.bbox.x1)

    margin_left = _percentile(lefts, 0.1) if lefts else _FALLBACK_MARGIN_LEFT_PT
    margin_right = (width_pt - _percentile(rights, 0.9)) if rights else _FALLBACK_MARGIN_RIGHT_PT
    margin_top = min(tops) if tops else _FALLBACK_MARGIN_TOP_PT
    margin_bottom = (height_pt - max(bottoms)) if bottoms else _FALLBACK_MARGIN_BOTTOM_PT
    # Страницы разделяются явными разрывами, поэтому нижнее поле не влияет
    # на вёрстку — но если содержимое в Word окажется чуть выше исходного,
    # большое нижнее поле вытолкнет последние строки на лишнюю страницу.
    if not document.footer:
        margin_bottom = min(margin_bottom, 20.0)

    def clamp(v: float, lo: float, hi: float, fallback: float) -> float:
        return v if lo <= v <= hi else fallback

    return PageGeometry(
        width_pt=width_pt,
        height_pt=height_pt,
        margin_left_pt=clamp(margin_left, 14.0, 140.0, _FALLBACK_MARGIN_LEFT_PT),
        margin_right_pt=clamp(margin_right, 14.0, 140.0, _FALLBACK_MARGIN_RIGHT_PT),
        margin_top_pt=clamp(margin_top, 10.0, 120.0, _FALLBACK_MARGIN_TOP_PT),
        margin_bottom_pt=clamp(margin_bottom, 10.0, 120.0, _FALLBACK_MARGIN_BOTTOM_PT),
    )


# --------------------------------------------------------------------------
# Шрифты и стили
# --------------------------------------------------------------------------


def _set_rfonts(rpr, font_name: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)


def _setup_styles(doc: DocxDocument, font_name: str) -> None:
    """Приводит стили шаблона python-docx к нейтральному виду: без
    тематических шрифтов (Calibri/Cambria), без синего цвета и собственных
    отступов у заголовков, без "space after 10pt / 1.15" у обычного
    текста. Иначе стиль заголовка/списка сам по себе меняет внешний вид
    абзаца относительно оригинала (синий крупный заголовок на месте
    обычной полужирной строки)."""
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is not None:
                _set_rfonts(rpr, font_name)
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    _set_rfonts(normal.element.get_or_add_rPr(), font_name)
    pf = normal.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for style in doc.styles:
        name = getattr(style, "name", "") or ""
        if not (name.startswith("Heading") or name in ("Title", "Subtitle") or name.startswith("List")):
            continue
        try:
            font = style.font
        except AttributeError:
            continue
        rpr = style.element.get_or_add_rPr()
        _set_rfonts(rpr, font_name)
        font.color.rgb = RGBColor(0, 0, 0)
        font.bold = None
        font.italic = None
        font.size = None
        if not hasattr(style, "paragraph_format"):
            continue  # символьный стиль ("Heading 1 Char")
        spf = style.paragraph_format
        # "Не добавлять интервал между абзацами одного стиля" (стили списков
        # шаблона) съедает измеренные отступы между пунктами списка.
        ppr = style.element.pPr
        if ppr is not None:
            for el in ppr.findall(qn("w:contextualSpacing")):
                ppr.remove(el)
        spf.space_before = Pt(0)
        spf.space_after = Pt(0)
        spf.keep_with_next = None
        spf.keep_together = None


def _apply_run_font(run, run_data, default_font: str) -> None:
    font_name = clean_font_name(run_data.font_name) or default_font
    run.font.name = font_name
    _set_rfonts(run._r.get_or_add_rPr(), font_name)


def _set_paragraph_mark_font(p, size_pt: float, font_name: str, bold: bool = False) -> None:
    """Формат "знака абзаца" — от него Word берёт вид номера пункта списка и
    высоту пустой строки."""
    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    _set_rfonts(rpr, font_name)
    if bold:
        rpr.append(OxmlElement("w:b"))
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(int(round(size_pt * 2))))
        rpr.append(el)


# --------------------------------------------------------------------------
# Поток блоков
# --------------------------------------------------------------------------


class _Box:
    """Горизонтальные границы контейнера (страница или ячейка) в
    координатах исходной страницы."""

    def __init__(self, left: float, right: float, max_outdent: float = 0.0, page_width: float = 0.0):
        self.left = left
        self.right = right
        self.page_width = page_width
        # Насколько абзац может выступать влево за границу контейнера
        # (на странице — в поле; в ячейке — нисколько).
        self.max_outdent = max_outdent


class _Flow:
    """Вертикальный "курсор" потока блоков внутри контейнера: y (в
    координатах исходной страницы), до которой уже дошёл текст в Word."""

    def __init__(self, top: Optional[float]):
        self.cursor = top
        self.last_paragraph = None
        self.last_was_table = False
        self.first_paragraph = None


class _Writer:
    def __init__(self, doc: DocxDocument, geometry: PageGeometry, pitch_ratio: float, font_name: str):
        self.doc = doc
        self.geometry = geometry
        self.pitch_ratio = pitch_ratio
        self.font_name = font_name
        self._docpr_id = 1000

    # --- абзацы -----------------------------------------------------------

    def _new_paragraph(self, container, flow: _Flow, style: Optional[str] = None):
        if flow.first_paragraph is None and hasattr(container, "_tc"):
            # В новой ячейке уже есть один пустой абзац — используем его.
            p = container.paragraphs[0]
            if style:
                p.style = style
        else:
            p = container.add_paragraph(style=style)
        if flow.first_paragraph is None:
            flow.first_paragraph = p
        flow.last_paragraph = p
        flow.last_was_table = False
        return p

    def write_paragraph(self, container, block: Paragraph, box: _Box, flow: _Flow, style: Optional[str] = None):
        p = self._new_paragraph(container, flow, style)
        pf = p.paragraph_format
        p.alignment = _ALIGN_MAP.get(block.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        pf.space_after = Pt(0)
        size = _dominant_size(block)

        if block.baseline_pt is not None:
            # Геометрия известна: точный межстрочный шаг и положение
            # базовой линии как в оригинале.
            pitch = _block_pitch(block, self.pitch_ratio)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(round(pitch, 2))
            box_top = block.baseline_pt - pitch * _BASELINE_RATIO
            before = 0.0
            if flow.cursor is not None:
                before = max(0.0, box_top - flow.cursor)
            pf.space_before = Pt(round(min(before, 400.0), 2))
            top = box_top if flow.cursor is None else max(flow.cursor, box_top)
            flow.cursor = top + pitch * max(1, block.n_lines)

            left = block.indent_pt - box.left
            if abs(left) < 1.0:
                left = 0.0
            left = max(left, -box.max_outdent)
            pf.left_indent = Pt(round(left, 2))
            if abs(block.first_line_indent_pt) >= 1.0:
                pf.first_line_indent = Pt(round(block.first_line_indent_pt, 2))
            if block.right_edge_pt is not None:
                right = box.right - block.right_edge_pt - _WRAP_SLACK_PT
                if right >= 1.0:
                    pf.right_indent = Pt(round(right, 2))
        else:
            # Блок без геометрии (например, заглушка ошибки страницы):
            # только измеренный зазор и одинарный интервал.
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_before = Pt(min(block.space_before_pt, 200.0)) if block.space_before_pt > 0.5 else Pt(0)
            # indent_pt — абсолютная координата на исходной странице; поле
            # left_indent в DOCX откладывается ДОПОЛНИТЕЛЬНО от левой
            # границы контейнера, поэтому нужна именно разница.
            extra_indent = block.indent_pt - box.left
            if extra_indent > 3.0:
                pf.left_indent = Pt(min(extra_indent, (box.right - box.left) * 0.6))
            if flow.cursor is not None and block.bbox is not None:
                flow.cursor = block.bbox.y1

        for run_data in block.runs:
            if not run_data.text:
                continue
            r = p.add_run(run_data.text)
            r.bold = run_data.bold
            r.italic = run_data.italic
            r.font.size = Pt(max(4.0, min(run_data.size_pt, 72.0)))
            _apply_run_font(r, run_data, self.font_name)
            if run_data.char_spacing_pt:
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:val"), str(int(round(run_data.char_spacing_pt * 20))))
                r._r.get_or_add_rPr().append(spacing)
        first_run = next((r for r in block.runs if r.text.strip()), None)
        mark_font = (clean_font_name(first_run.font_name) if first_run else None) or self.font_name
        _set_paragraph_mark_font(p, size, mark_font, bold=bool(first_run and first_run.bold and isinstance(block, ListItem)))
        return p

    def _spacer(self, container, flow: _Flow, height: float):
        """Пустой абзац заданной высоты (перед таблицей, между таблицами)."""
        p = self._new_paragraph(container, flow)
        pf = p.paragraph_format
        h = max(1.0, height)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(round(h, 2))
        _set_paragraph_mark_font(p, min(h, 11.0), self.font_name)
        if flow.cursor is not None:
            flow.cursor += h
        return p

    # --- таблицы ----------------------------------------------------------

    def write_table(self, container, table: Table, box: _Box, flow: _Flow) -> None:
        n_rows = table.n_rows
        n_cols = table.n_cols
        if n_rows == 0 or n_cols == 0:
            return

        # Верх таблицы-раскладки (колонки текста без рамок) — верх первой
        # строки текста в ней, а не верх "чернил" этой строки: иначе текст
        # колонок оказывается ниже, чем в оригинале.
        table_top = table.bbox.y0 if table.bbox is not None else None
        if table.borderless and table_top is not None:
            for row in table.rows:
                for cell in row.cells:
                    first = next((b for b in cell.blocks if isinstance(b, Paragraph) and b.text.strip()), None)
                    box_top = _block_box_top(first, self.pitch_ratio) if first is not None else None
                    if box_top is not None:
                        table_top = min(table_top, box_top)

        # Зазор перед таблицей: у таблицы нет "отступа перед", поэтому он
        # переносится в "отступ после" предыдущего абзаца (или в пустой
        # абзац нужной высоты, если перед таблицей нет абзаца/стоит другая
        # таблица — Word склеил бы две соседние таблицы в одну).
        if table_top is not None and flow.cursor is not None:
            gap = table_top - flow.cursor
            if flow.last_paragraph is not None and not flow.last_was_table:
                if gap > 0.5:
                    flow.last_paragraph.paragraph_format.space_after = Pt(round(min(gap, 400.0), 2))
            elif gap > 0.5 or flow.last_was_table:
                self._spacer(container, flow, gap)
        elif flow.last_was_table:
            self._spacer(container, flow, 1.0)

        docx_table = container.add_table(rows=n_rows, cols=n_cols)
        if not table.borderless:
            docx_table.style = "Table Grid"
        docx_table.autofit = False
        self._set_table_layout(docx_table, table, box)

        # Таблица может выходить за поля текста (как и в оригинале) —
        # ограничиваем ширину только краем листа, а не полями: сжатие
        # столбцов под поля ломало слова посередине ("Алексееви-ч").
        usable = box.right - box.left
        if table.bbox is not None and box.page_width:
            usable = max(usable, box.page_width - 5.0 - max(5.0, table.bbox.x0))
        col_widths_pt = _fitted_column_widths(table.col_widths_pt, n_cols, usable)

        # Сначала объединяем ячейки, потом заполняем: python-docx при
        # объединении переносит содержимое поглощаемых ячеек в итоговую.
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if cell.is_merge_continuation:
                    continue
                if cell.row_span > 1 or cell.col_span > 1:
                    r1, c1 = r + cell.row_span - 1, c + cell.col_span - 1
                    try:
                        docx_table.cell(r, c).merge(docx_table.cell(r1, c1))
                    except IndexError:
                        logger.warning("Не удалось объединить ячейки таблицы (%d,%d)-(%d,%d)", r, c, r1, c1)

        for c in range(n_cols):
            docx_table.columns[c].width = Pt(col_widths_pt[c])

        padding = _table_padding(table)
        col_x = [table.bbox.x0 if table.bbox else box.left]
        for w in col_widths_pt:
            col_x.append(col_x[-1] + w)

        for r, row in enumerate(table.rows):
            docx_row = docx_table.rows[r]
            if row.height_pt and row.height_pt > 2:
                docx_row.height = Pt(round(row.height_pt, 2))
                docx_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for c, cell in enumerate(row.cells):
                if cell.is_merge_continuation:
                    continue
                docx_cell = docx_table.cell(r, c)
                span_w = sum(col_widths_pt[c : c + max(1, cell.col_span)])
                docx_cell.width = Pt(span_w)
                docx_cell.vertical_alignment = _VALIGN_MAP.get(cell.v_align, WD_CELL_VERTICAL_ALIGNMENT.TOP)
                if cell.bbox is not None:
                    cell_pad = _cell_padding(cell, padding)
                    if cell_pad < padding - 0.05:
                        _set_cell_margins(docx_cell, cell_pad)
                    cell_box = _Box(cell.bbox.x0 + cell_pad, cell.bbox.x1 - cell_pad)
                    cell_top = table_top if table.borderless and table_top is not None else cell.bbox.y0
                    cell_flow = _Flow(cell_top if cell.v_align == "top" else None)
                else:
                    x0 = col_x[c] + padding
                    cell_box = _Box(x0, x0 + span_w - 2 * padding)
                    cell_flow = _Flow(None)
                wrote_any = False
                for block in cell.blocks:
                    if isinstance(block, Paragraph) and not block.text.strip():
                        continue
                    self.write_block(docx_cell, block, cell_box, cell_flow)
                    wrote_any = True
                if not wrote_any:
                    # Пустая ячейка: высота её единственного пустого абзаца
                    # должна соответствовать шрифту таблицы, а не 11pt.
                    size = next(
                        (_dominant_size(b) for b in cell.blocks if isinstance(b, Paragraph) and b.runs), 9.0
                    )
                    p = docx_cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    _set_paragraph_mark_font(p, size, self.font_name)

        flow.last_paragraph = None
        flow.last_was_table = True
        if table.bbox is not None and flow.cursor is not None:
            flow.cursor = max(flow.cursor, table.bbox.y1)

    def _set_table_layout(self, docx_table, table: Table, box: _Box) -> None:
        tbl_pr = docx_table._tbl.tblPr
        # Фиксированная раскладка: ширины столбцов как в оригинале, Word
        # не "подгоняет" их под содержимое.
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
        margin = _table_padding(table)
        mar = OxmlElement("w:tblCellMar")
        for side, value in (("top", 0.0), ("left", margin), ("bottom", 0.0), ("right", margin)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(int(round(value * 20))))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tbl_pr.append(mar)
        if table.bbox is not None:
            # В режиме совместимости шаблона python-docx (Word 2010) отступ
            # таблицы отсчитывается до текста первой ячейки, а не до её
            # границы — добавляем внутреннее поле, чтобы граница таблицы
            # оказалась там же, где в оригинале.
            indent = table.bbox.x0 - box.left + margin
            if abs(indent) >= 1.0:
                ind = OxmlElement("w:tblInd")
                ind.set(qn("w:w"), str(int(round(indent * 20))))
                ind.set(qn("w:type"), "dxa")
                tbl_pr.append(ind)

    # --- изображения ------------------------------------------------------

    def write_inline_image(self, container, block: ImageBlock, box: _Box, flow: _Flow) -> None:
        try:
            pil_img = PILImage.open(io.BytesIO(block.data))
            width_px, height_px = pil_img.size
        except Exception:  # noqa: BLE001
            width_px, height_px = block.width_px or 100, block.height_px or 100

        usable = box.right - box.left
        # Реальный размер изображения на исходной странице (если известен из
        # bbox) — не растягиваем логотип/печать на всю ширину страницы только
        # потому, что так проще; иначе маленькая картинка становится нелепо
        # большой и разваливает окружающую вёрстку.
        if block.bbox is not None and (block.bbox.x1 - block.bbox.x0) > 1:
            width_pt = min(block.bbox.x1 - block.bbox.x0, usable)
        else:
            px_per_pt = 96.0 / 72.0  # обычный экранный масштаб как разумная оценка при отсутствии bbox
            width_pt = min(width_px / px_per_pt, usable)
        width_pt = max(width_pt, 36.0)

        p = self._new_paragraph(container, flow)
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if block.bbox is not None:
            if flow.cursor is not None:
                pf.space_before = Pt(round(max(0.0, min(block.bbox.y0 - flow.cursor, 400.0)), 2))
            indent = block.bbox.x0 - box.left
            if indent > 1.0 and indent + width_pt <= usable:
                pf.left_indent = Pt(round(indent, 2))
        run = p.add_run()
        try:
            # width= достаточно: python-docx сохраняет пропорции изображения
            # автоматически, когда передан только один из width/height.
            run.add_picture(io.BytesIO(block.data), width=Pt(width_pt))
        except Exception as exc:  # noqa: BLE001
            logger.warning("Не удалось вставить изображение: %s", exc)
            return
        if flow.cursor is not None:
            height_pt = width_pt * height_px / max(width_px, 1)
            flow.cursor = max(flow.cursor, block.bbox.y0 if block.bbox else flow.cursor) + height_pt
        if block.caption:
            cap = self._new_paragraph(container, flow)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(0)
            run = cap.add_run(block.caption)
            run.italic = True
            run.font.size = Pt(9)
            if flow.cursor is not None:
                flow.cursor += 11.0

    def attach_floating_image(self, paragraph, block: ImageBlock) -> None:
        """Плавающее изображение за текстом в точных координатах страницы."""
        if block.bbox is None:
            return
        width_pt = max(1.0, block.bbox.x1 - block.bbox.x0)
        height_pt = max(1.0, block.bbox.y1 - block.bbox.y0)
        run = paragraph.add_run()
        try:
            inline_shape = run.add_picture(io.BytesIO(block.data), width=Pt(width_pt), height=Pt(height_pt))
        except Exception as exc:  # noqa: BLE001
            logger.warning("Не удалось вставить изображение: %s", exc)
            return
        inline = inline_shape._inline
        self._docpr_id += 1
        anchor = OxmlElement("wp:anchor")
        for k, v in (
            ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"),
            ("relativeHeight", str(self._docpr_id)), ("behindDoc", "1"), ("locked", "0"),
            ("layoutInCell", "1"), ("allowOverlap", "1"),
        ):
            anchor.set(k, v)
        simple = OxmlElement("wp:simplePos")
        simple.set("x", "0")
        simple.set("y", "0")
        anchor.append(simple)
        for tag, rel, value in (("wp:positionH", "page", block.bbox.x0), ("wp:positionV", "page", block.bbox.y0)):
            pos = OxmlElement(tag)
            pos.set("relativeFrom", rel)
            off = OxmlElement("wp:posOffset")
            off.text = str(int(Emu(Pt(value))))
            pos.append(off)
            anchor.append(pos)
        anchor.append(copy.deepcopy(inline.find(qn("wp:extent"))))
        effect = OxmlElement("wp:effectExtent")
        for k in ("l", "t", "r", "b"):
            effect.set(k, "0")
        anchor.append(effect)
        anchor.append(OxmlElement("wp:wrapNone"))
        docpr = copy.deepcopy(inline.find(qn("wp:docPr")))
        docpr.set("id", str(self._docpr_id))
        anchor.append(docpr)
        anchor.append(OxmlElement("wp:cNvGraphicFramePr"))
        anchor.append(copy.deepcopy(inline.find(qn("a:graphic"))))
        inline.getparent().replace(inline, anchor)

    # --- диспетчер ----------------------------------------------------------

    def write_block(self, container, block, box: _Box, flow: _Flow) -> None:
        if isinstance(block, Heading):
            self.write_paragraph(container, block, box, flow, style=f"Heading {min(max(block.level, 1), 4)}")
        elif isinstance(block, ListItem):
            self.write_paragraph(container, block, box, flow, style=_list_style(block))
        elif isinstance(block, Table):
            self.write_table(container, block, box, flow)
        elif isinstance(block, ImageBlock):
            self.write_inline_image(container, block, box, flow)
        elif isinstance(block, Paragraph):
            if block.text.strip():
                self.write_paragraph(container, block, box, flow)


def _fitted_column_widths(col_widths_pt: list[float], n_cols: int, usable_width_pt: float) -> list[float]:
    """Пропорционально вписывает известные ширины столбцов в печатную
    область страницы (после того как страница DOCX приведена к размеру
    оригинала, они обычно и так укладываются, но подстраховываемся)."""
    if not col_widths_pt or len(col_widths_pt) != n_cols:
        even = usable_width_pt / n_cols
        return [even] * n_cols
    total = sum(col_widths_pt) or 1.0
    if total <= usable_width_pt:
        return list(col_widths_pt)
    scale = usable_width_pt / total
    return [w * scale for w in col_widths_pt]


# Стандартный шаблон Word (тот, что использует python-docx по умолчанию)
# не даёт единого многоуровневого списка через один стиль — вместо этого
# для второго/третьего уровня вложенности есть отдельные стили с уже
# готовой (более глубокой) нумерацией/отступом ("List Number 2/3",
# "List Bullet 2/3"). Больше трёх уровней в исходном шаблоне не
# предусмотрено — упираемся в третий и полагаемся на left_indent
# (см. _Writer.write_paragraph) для более глубокой вложенности.
_LIST_STYLES_ORDERED = ["List Number", "List Number 2", "List Number 3"]
_LIST_STYLES_BULLET = ["List Bullet", "List Bullet 2", "List Bullet 3"]


def _list_style(block: ListItem) -> str:
    styles = _LIST_STYLES_ORDERED if block.ordered else _LIST_STYLES_BULLET
    idx = min(max(block.level, 0), len(styles) - 1)
    return styles[idx]


def _cell_padding(cell, table_padding: float) -> float:
    """Поле конкретной ячейки: если самая широкая строка ячейки в оригинале
    почти касается границ, общего поля таблицы ей не хватит — строка в Word
    перенеслась бы посреди слова."""
    widths = [b.bbox.x1 - b.bbox.x0 for b in cell.blocks if isinstance(b, Paragraph) and b.bbox is not None]
    if not widths or cell.bbox is None:
        return table_padding
    free = (cell.bbox.x1 - cell.bbox.x0 - max(widths)) / 2 - 1.0
    return max(0.0, min(table_padding, free))


def _set_cell_margins(docx_cell, margin_pt: float) -> None:
    tc_pr = docx_cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(round(margin_pt * 20))))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _table_padding(table: Table) -> float:
    if table.borderless:
        return 0.0
    if table.cell_padding_pt is not None:
        return table.cell_padding_pt
    return _CELL_MARGIN_PT


def _document_pitch_ratio(document: Document) -> float:
    """Типичное отношение межстрочного шага к кеглю в документе — для
    однострочных абзацев, у которых шаг не измерить напрямую."""
    ratios = []
    for page in document.pages:
        for b in page.blocks:
            if isinstance(b, Paragraph) and b.line_pitch_pt and b.n_lines >= 2:
                size = _dominant_size(b)
                if size > 0:
                    ratios.append(b.line_pitch_pt / size)
    ratios = [r for r in ratios if 0.9 <= r <= 2.5]
    return statistics.median(ratios) if ratios else _DEFAULT_PITCH_RATIO


def build_docx(document: Document, output_path: str, default_font: str = DEFAULT_FONT_NAME) -> None:
    try:
        doc = DocxDocument()
        _setup_styles(doc, default_font)
        pitch_ratio = _document_pitch_ratio(document)
        geometry = _compute_page_geometry(document, pitch_ratio)

        section = doc.sections[0]
        section.page_width = Pt(geometry.width_pt)
        section.page_height = Pt(geometry.height_pt)
        section.left_margin = Pt(geometry.margin_left_pt)
        section.right_margin = Pt(geometry.margin_right_pt)
        section.top_margin = Pt(geometry.margin_top_pt)
        section.bottom_margin = Pt(geometry.margin_bottom_pt)
        section.header_distance = Pt(min(geometry.margin_top_pt, 20.0))
        section.footer_distance = Pt(min(geometry.margin_bottom_pt, 20.0))

        if document.footer and document.footer.text:
            footer_p = section.footer.paragraphs[0]
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_p.add_run(document.footer.text + "  ")
            _add_field(footer_p, "PAGE")

        if document.header and document.header.text:
            header_p = section.header.paragraphs[0]
            header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header_p.add_run(document.header.text)

        writer = _Writer(doc, geometry, pitch_ratio, default_font)
        page_box = _Box(
            geometry.margin_left_pt,
            geometry.width_pt - geometry.margin_right_pt,
            max_outdent=max(0.0, geometry.margin_left_pt - 5.0),
            page_width=geometry.width_pt,
        )
        for page_index, page in enumerate(document.pages):
            flow = _Flow(geometry.margin_top_pt)
            floating = [b for b in page.blocks if isinstance(b, ImageBlock) and b.floating and b.bbox]
            flow_blocks = [b for b in page.blocks if not (isinstance(b, ImageBlock) and b.floating and b.bbox)]
            if page_index > 0 or floating:
                # Абзац-"якорь" страницы нужен для разрыва страницы перед ней
                # (разрыв как свойство абзаца, а не отдельный абзац с
                # символом разрыва — тот сдвигал бы всю следующую страницу на
                # строку вниз) и для привязки плавающих изображений.
                first = flow_blocks[0] if flow_blocks else None
                if not (isinstance(first, Paragraph) and first.text.strip()):
                    writer._spacer(doc, flow, 1.0)
            for block in flow_blocks:
                writer.write_block(doc, block, page_box, flow)
            anchor = flow.first_paragraph
            if anchor is None:
                anchor = writer._spacer(doc, flow, 1.0)
            if page_index > 0:
                anchor.paragraph_format.page_break_before = True
            for img in floating:
                writer.attach_floating_image(anchor, img)

        # Документ не должен заканчиваться таблицей (Word требует абзац
        # последним элементом тела документа).
        body = doc.element.body
        last = [el for el in body if el.tag != qn("w:sectPr")]
        if last and last[-1].tag == qn("w:tbl"):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(1)
            _set_paragraph_mark_font(p, 1.0, default_font)

        doc.save(output_path)
    except Exception as exc:  # noqa: BLE001
        raise DocxWriteError(f"Не удалось сохранить DOCX: {exc}") from exc
