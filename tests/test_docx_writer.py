import os

from docx import Document as DocxDocument

from bes_ocr.core.docx_writer import (
    _compute_page_geometry,
    _fitted_column_widths,
    _list_style,
    build_docx,
)
from bes_ocr.core.models import (
    BBox,
    Document,
    Heading,
    ImageBlock,
    ListItem,
    Page,
    Paragraph,
    Run,
    Table,
    TableCell,
    TableRow,
)


def _sample_document() -> Document:
    doc = Document()
    page = Page(number=1)
    page.blocks.append(Heading(runs=[Run(text="Title")], level=1))
    page.blocks.append(Paragraph(runs=[Run(text="Hello world")]))

    table = Table()
    row0 = TableRow(cells=[TableCell(blocks=[Paragraph(runs=[Run(text="Merged")])], row_span=1, col_span=2), TableCell(is_merge_continuation=True, row_span=0, col_span=0)])
    row1 = TableRow(cells=[TableCell(blocks=[Paragraph(runs=[Run(text="A")])]), TableCell(blocks=[Paragraph(runs=[Run(text="B")])])])
    table.rows = [row0, row1]
    page.blocks.append(table)

    doc.pages.append(page)
    return doc


def test_build_docx_creates_valid_file(tmp_path):
    document = _sample_document()
    out_path = os.path.join(tmp_path, "out.docx")
    build_docx(document, out_path)
    assert os.path.isfile(out_path)

    reopened = DocxDocument(out_path)
    assert any(p.text == "Title" and p.style.name.startswith("Heading") for p in reopened.paragraphs)
    assert any(p.text == "Hello world" for p in reopened.paragraphs)
    assert len(reopened.tables) == 1
    t = reopened.tables[0]
    assert len(t.rows) == 2
    assert t.cell(0, 0).text == "Merged"
    assert t.cell(0, 1).text == "Merged"  # объединённая ячейка
    assert t.cell(1, 0).text == "A"
    assert t.cell(1, 1).text == "B"


def test_build_docx_with_image(tmp_path):
    from PIL import Image
    import io

    img = Image.new("RGB", (50, 50), "red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    document = Document()
    page = Page(number=1)
    page.blocks.append(ImageBlock(data=buf.getvalue(), width_px=50, height_px=50, caption="A caption"))
    document.pages.append(page)

    out_path = os.path.join(tmp_path, "img.docx")
    build_docx(document, out_path)

    reopened = DocxDocument(out_path)
    assert len(reopened.inline_shapes) == 1


def test_page_geometry_matches_source_pdf_page_size(tmp_path):
    """DOCX должен получать размер страницы исходного PDF, а не Letter по
    умолчанию — иначе перенос строк и пропорции таблиц систематически
    расходятся с оригиналом даже при верно распознанном содержимом."""
    document = _sample_document()
    document.pages[0].width_pt = 595.0
    document.pages[0].height_pt = 842.0

    out_path = os.path.join(tmp_path, "geom.docx")
    build_docx(document, out_path)

    reopened = DocxDocument(out_path)
    section = reopened.sections[0]
    assert round(section.page_width.pt) == 595
    assert round(section.page_height.pt) == 842


def test_compute_page_geometry_margins_from_content():
    document = _sample_document()
    document.pages[0].width_pt = 595.0
    document.pages[0].height_pt = 842.0
    for b in document.pages[0].blocks:
        if isinstance(b, (Heading, Paragraph)):
            b.bbox = BBox(60.0, 40.0, 500.0, 60.0)

    geometry = _compute_page_geometry(document)
    assert 40 < geometry.margin_left_pt < 90
    assert geometry.usable_width_pt < geometry.width_pt


def test_compute_page_geometry_falls_back_without_pages():
    geometry = _compute_page_geometry(Document())
    assert geometry.width_pt > 0
    assert geometry.margin_left_pt > 0


def test_fitted_column_widths_scales_down_when_too_wide():
    widths = _fitted_column_widths([100.0, 100.0, 100.0], 3, usable_width_pt=150.0)
    assert sum(widths) <= 150.0 + 1e-6
    # пропорции столбцов сохранены (все три равны)
    assert widths[0] == widths[1] == widths[2]


def test_fitted_column_widths_falls_back_to_even_split():
    widths = _fitted_column_widths([], 4, usable_width_pt=400.0)
    assert widths == [100.0, 100.0, 100.0, 100.0]


def test_table_column_widths_written_to_docx(tmp_path):
    document = Document()
    page = Page(number=1, width_pt=595.0, height_pt=842.0)
    table = Table(col_widths_pt=[200.0, 100.0])
    table.rows = [
        TableRow(
            cells=[
                TableCell(blocks=[Paragraph(runs=[Run(text="Wide")])]),
                TableCell(blocks=[Paragraph(runs=[Run(text="Narrow")])]),
            ]
        )
    ]
    page.blocks.append(table)
    document.pages.append(page)

    out_path = os.path.join(tmp_path, "widths.docx")
    build_docx(document, out_path)

    reopened = DocxDocument(out_path)
    col_widths = [c.width.pt for c in reopened.tables[0].columns]
    assert col_widths[0] > col_widths[1]


def test_table_cell_font_size_matches_run_size_pt(tmp_path):
    """Регрессия: раньше текст ячеек всегда получал шрифт стиля Word по
    умолчанию (~11pt) независимо от `Run.size_pt`, из-за чего плотные
    таблицы реальных бланков (распознанный шрифт которых мельче) переносили
    заметно больше строк, чем в оригинале, и документ раздувался на лишние
    страницы (измерено на реальном документе: 2 страницы → 3)."""
    document = Document()
    page = Page(number=1, width_pt=595.0, height_pt=842.0)
    table = Table(col_widths_pt=[200.0])
    table.rows = [
        TableRow(cells=[TableCell(blocks=[Paragraph(runs=[Run(text="Small", size_pt=7.0)])])])
    ]
    page.blocks.append(table)
    document.pages.append(page)

    out_path = os.path.join(tmp_path, "cell_font.docx")
    build_docx(document, out_path)

    reopened = DocxDocument(out_path)
    # cell.text = "" (см. docx_writer._write_table) уже оставляет один пустой
    # run в параграфе, наш форматированный текст добавляется следующим.
    run = reopened.tables[0].cell(0, 0).paragraphs[0].runs[-1]
    assert run.font.size.pt == 7.0


def test_list_style_picks_nested_style_by_level():
    assert _list_style(ListItem(ordered=True, level=0)) == "List Number"
    assert _list_style(ListItem(ordered=True, level=1)) == "List Number 2"
    assert _list_style(ListItem(ordered=True, level=2)) == "List Number 3"
    # За пределами глубины, предусмотренной шаблоном Word, остаёмся на
    # самом глубоком доступном стиле, а не падаем/выходим за границы списка.
    assert _list_style(ListItem(ordered=True, level=5)) == "List Number 3"
    assert _list_style(ListItem(ordered=False, level=1)) == "List Bullet 2"


def test_nested_list_items_get_distinct_docx_styles(tmp_path):
    document = Document()
    page = Page(number=1, width_pt=595.0, height_pt=842.0)
    page.blocks.append(ListItem(runs=[Run(text="Top")], ordered=True, level=0))
    page.blocks.append(ListItem(runs=[Run(text="Nested")], ordered=True, level=1))
    document.pages.append(page)

    out_path = os.path.join(tmp_path, "nested_list.docx")
    build_docx(document, out_path)

    reopened = DocxDocument(out_path)
    styles = [p.style.name for p in reopened.paragraphs if p.text in ("Top", "Nested")]
    assert styles == ["List Number", "List Number 2"]
